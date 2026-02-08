import os

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


def ensure_sample_workbook(path: str) -> str:
    if os.path.exists(path):
        return path
    os.makedirs(os.path.dirname(path), exist_ok=True)
    create_sample_workbook(path)
    return path


def create_sample_workbook(path: str) -> None:
    workbook = Workbook()

    ws_inputs = workbook.active
    ws_inputs.title = "Inputs"
    ws_inputs["A1"] = "key"
    ws_inputs["B1"] = "value"
    ws_inputs["A2"] = "BaseRate"
    ws_inputs["B2"] = 0.05
    ws_inputs["A3"] = "Volume"
    ws_inputs["B3"] = 1200
    ws_inputs["A4"] = "Adjustment"
    ws_inputs["B4"] = 1.1

    ws_tables = workbook.create_sheet("RateTable")
    ws_tables["A1"] = "Tier"
    ws_tables["B1"] = "Rate"
    ws_tables["A2"] = 1
    ws_tables["B2"] = 0.02
    ws_tables["A3"] = 2
    ws_tables["B3"] = 0.03
    ws_tables["A4"] = 3
    ws_tables["B4"] = 0.04

    ws_formulas = workbook.create_sheet("Formulas")
    ws_formulas["A1"] = "name"
    ws_formulas["B1"] = "formula"
    ws_formulas["A2"] = "AdjustedRate"
    ws_formulas["B2"] = "=BaseRate+Adjustment"
    ws_formulas["A3"] = "VolumeCharge"
    ws_formulas["B3"] = "=Volume*BaseRate"
    ws_formulas["A4"] = "TableAdjustedCharge"
    ws_formulas["B4"] = "=Volume*BaseRate*RateTable"

    workbook.defined_names.append(
        DefinedName("BaseRate", attr_text="Inputs!$B$2")
    )
    workbook.defined_names.append(
        DefinedName("Volume", attr_text="Inputs!$B$3")
    )
    workbook.defined_names.append(
        DefinedName("Adjustment", attr_text="Inputs!$B$4")
    )
    workbook.defined_names.append(
        DefinedName("RateTable", attr_text="RateTable!$A$1:$B$4")
    )

    workbook.save(path)
