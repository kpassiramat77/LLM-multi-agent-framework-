from __future__ import annotations

from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


def extract_workbook_data(path: str) -> Dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    factors = extract_factors(workbook)
    tables = extract_tables(workbook)
    formulas = extract_formulas(workbook)
    return {
        "factors": factors,
        "tables": tables,
        "formulas": formulas,
    }


def extract_factors(workbook) -> List[Dict[str, Any]]:
    factors: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        if is_factor_sheet(sheet):
            start_row = 1
            header = [cell.value for cell in sheet[1][:2]]
            if header and looks_like_factor_header(header):
                start_row = 2
            for row in sheet.iter_rows(min_row=start_row):
                key_cell = row[0].value if len(row) > 0 else None
                value_cell = row[1].value if len(row) > 1 else None
                key = normalize_key(key_cell)
                if not key:
                    continue
                factors.append(
                    {
                        "name": key,
                        "value": normalize_value(value_cell),
                    }
                )
            break
    return factors


def extract_tables(workbook) -> List[Dict[str, Any]]:
    tables: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        if is_factor_sheet(sheet):
            continue
        table = extract_single_table_from_sheet(sheet)
        if table:
            tables.append(table)
    return tables


def extract_formulas(workbook) -> List[Dict[str, Any]]:
    formulas: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    label = None
                    if cell.column > 1:
                        left = sheet.cell(row=cell.row, column=cell.column - 1).value
                        if isinstance(left, str) and left.strip():
                            label = left.strip()
                    formulas.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "formula": cell.value,
                            "label": label,
                        }
                    )
    return formulas


def is_factor_sheet(sheet) -> bool:
    title = (sheet.title or "").strip().lower()
    if title in {"inputs", "factors"}:
        return True
    header = [cell.value for cell in sheet[1][:2]]
    return looks_like_factor_header(header)


def looks_like_factor_header(header: List[Any]) -> bool:
    if len(header) < 2:
        return False
    first = normalize_key(header[0])
    second = normalize_key(header[1])
    return first in {"name", "key", "input"} and second in {"value", "val"}


def extract_single_table_from_sheet(sheet) -> Optional[Dict[str, Any]]:
    header_row = sheet[1]
    columns: List[str] = []
    col_indices: List[int] = []
    for idx, cell in enumerate(header_row, start=1):
        if cell.value is None:
            break
        if isinstance(cell.value, str) and cell.value.strip():
            columns.append(cell.value.strip())
            col_indices.append(idx)
        else:
            return None
    if len(columns) < 2:
        return None

    first_data_row = 2
    numeric_row = [
        sheet.cell(row=first_data_row, column=col).value for col in col_indices
    ]
    if not any(is_numeric(value) for value in numeric_row):
        return None

    rows: List[List[Optional[float]]] = []
    row_idx = first_data_row
    while True:
        row_values = [sheet.cell(row=row_idx, column=col).value for col in col_indices]
        if all(value is None for value in row_values):
            break
        row_items: List[Optional[float]] = []
        for value in row_values:
            if is_numeric(value):
                row_items.append(float(value))
            elif value is None:
                row_items.append(None)
            else:
                row_items.append(None)
        rows.append(row_items)
        row_idx += 1

    if not rows:
        return None

    return {
        "name": sheet.title,
        "columns": columns,
        "rows": rows,
    }


def normalize_key(value: Any) -> Optional[str]:
    if isinstance(value, str):
        key = value.strip()
        if key:
            return key
    return None


def normalize_value(value: Any) -> Any:
    if isinstance(value, bool):
        return value
    if is_numeric(value):
        return float(value)
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    return str(value)


def is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)
