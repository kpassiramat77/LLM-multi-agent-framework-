import csv
import os
import re
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from prompt_lab.utils import ensure_dir


def convert_workbook_to_csvs(workbook_path: str, output_dir: str) -> Dict[str, Any]:
    ensure_dir(output_dir)
    workbook = load_workbook(workbook_path, data_only=False)

    inputs_sheet = find_sheet_by_name(workbook, "Inputs")
    if inputs_sheet is None:
        raise ValueError("Workbook missing required sheet: Inputs")

    inputs_csv = os.path.join(output_dir, "inputs.csv")
    write_inputs_csv(inputs_sheet, inputs_csv)

    formula_sheet = find_formula_sheet(workbook)
    formulas = extract_formulas(workbook, formula_sheet)
    formulas_csv = os.path.join(output_dir, "formulas.csv")
    write_formulas_csv(formulas, formulas_csv)

    skip_sheets = {inputs_sheet.title}
    if formula_sheet:
        skip_sheets.add(formula_sheet.title)
    tables_manifest = write_table_csvs(workbook, output_dir, skip_sheets)

    return {
        "inputs_csv": inputs_csv,
        "formulas_csv": formulas_csv,
        "tables": tables_manifest,
    }


def find_sheet_by_name(workbook, name: str):
    for sheet in workbook.worksheets:
        if (sheet.title or "").strip().lower() == name.lower():
            return sheet
    return None


def find_formula_sheet(workbook):
    for sheet in workbook.worksheets:
        if (sheet.title or "").strip().lower() == "formulas":
            return sheet
    for sheet in workbook.worksheets:
        if looks_like_formula_header(sheet):
            return sheet
    return None


def looks_like_formula_header(sheet) -> bool:
    header = [cell.value for cell in sheet[1][:2]]
    if len(header) < 2:
        return False
    first = normalize_header(header[0])
    second = normalize_header(header[1])
    return first in {"name", "calculation"} and second in {"formula", "expression"}


def write_inputs_csv(sheet, path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["key", "value"])
        for row in sheet.iter_rows(min_row=2):
            key_cell = row[0].value if len(row) > 0 else None
            value_cell = row[1].value if len(row) > 1 else None
            if key_cell is None and value_cell is None:
                break
            key = normalize_key(key_cell)
            if not key:
                continue
            writer.writerow([key, format_scalar(value_cell)])


def write_formulas_csv(formulas: List[Dict[str, Any]], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["name", "formula", "source_sheet", "cell"])
        for formula in formulas:
            writer.writerow(
                [
                    formula.get("name", ""),
                    formula.get("formula", ""),
                    formula.get("source_sheet", ""),
                    formula.get("cell", ""),
                ]
            )


def write_table_csvs(workbook, output_dir: str, skip_sheets: set) -> List[Dict[str, Any]]:
    tables_manifest = []
    for sheet in workbook.worksheets:
        if sheet.title in skip_sheets:
            continue
        table = extract_table_data(sheet)
        if not table:
            continue
        filename = f"table_{slugify(sheet.title)}.csv"
        path = os.path.join(output_dir, filename)
        write_table_csv(table, path)
        tables_manifest.append({"name": table["name"], "path": path})
    return tables_manifest


def extract_table_data(sheet) -> Optional[Dict[str, Any]]:
    header_row = sheet[1]
    columns: List[str] = []
    col_indices: List[int] = []
    for idx, cell in enumerate(header_row, start=1):
        if cell.value is None:
            break
        if isinstance(cell.value, str) and cell.value.strip():
            columns.append(cell.value.strip())
            col_indices.append(idx)
        else:
            return None
    if len(columns) < 2:
        return None

    first_data_row = 2
    numeric_row = [
        sheet.cell(row=first_data_row, column=col).value for col in col_indices
    ]
    if not any(is_numeric(value) for value in numeric_row):
        return None

    rows: List[List[Optional[float]]] = []
    row_idx = first_data_row
    while True:
        row_values = [sheet.cell(row=row_idx, column=col).value for col in col_indices]
        if all(value is None for value in row_values):
            break
        row_items: List[Optional[float]] = []
        for value in row_values:
            if is_numeric(value):
                row_items.append(float(value))
            else:
                row_items.append(None)
        rows.append(row_items)
        row_idx += 1

    if not rows:
        return None

    return {
        "name": sheet.title,
        "columns": columns,
        "rows": rows,
    }


def write_table_csv(table: Dict[str, Any], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(table["columns"])
        for row in table["rows"]:
            writer.writerow([format_scalar(value) for value in row])


def extract_formulas(workbook, formula_sheet=None) -> List[Dict[str, Any]]:
    if formula_sheet:
        return extract_formulas_from_sheet(formula_sheet)
    return extract_formulas_from_workbook(workbook)


def extract_formulas_from_sheet(sheet) -> List[Dict[str, Any]]:
    formulas: List[Dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2):
        name = normalize_key(row[0].value) if len(row) > 0 else None
        formula_cell = row[1] if len(row) > 1 else None
        formula_value = formula_cell.value if formula_cell else None
        if formula_value is None:
            continue
        formula_text = str(formula_value).strip()
        if not formula_text.startswith("="):
            formula_text = "=" + formula_text
        if not name:
            name = f"{sheet.title}_{formula_cell.coordinate}"
        formulas.append(
            {
                "name": name,
                "formula": formula_text,
                "source_sheet": sheet.title,
                "cell": formula_cell.coordinate,
            }
        )
    return formulas


def extract_formulas_from_workbook(workbook) -> List[Dict[str, Any]]:
    formulas: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    name = None
                    if cell.column > 1:
                        left = sheet.cell(row=cell.row, column=cell.column - 1).value
                        if isinstance(left, str) and left.strip():
                            name = left.strip()
                    if not name:
                        name = f"{sheet.title}_{cell.coordinate}"
                    formulas.append(
                        {
                            "name": name,
                            "formula": cell.value,
                            "source_sheet": sheet.title,
                            "cell": cell.coordinate,
                        }
                    )
    return formulas


def normalize_header(value: Any) -> Optional[str]:
    if isinstance(value, str):
        text = value.strip().lower()
        if text:
            return text
    return None


def normalize_key(value: Any) -> Optional[str]:
    if isinstance(value, str):
        key = value.strip()
        if key:
            return key
    return None


def format_scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if is_numeric(value):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    return str(value)


def is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def slugify(text: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    cleaned = cleaned.strip("_")
    return cleaned.lower() or "table"
